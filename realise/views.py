import csv
import io
import json
import logging
import re
import time
from datetime import datetime

from django.http import JsonResponse, HttpResponse
from django.shortcuts import redirect, render
from django.views.decorators.http import require_http_methods
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.cache import never_cache

from core.decorators import group_required, permission_flag_required, any_permission_flag
from core import sap_connector
from core.simple_xlsx import build_workbook
from . import services

logger = logging.getLogger(__name__)

REALISE_GROUPS = ('realise_admin', 'realise_premium', 'realise_commodity')

# In-memory cache: stores the last fetched raw SAP rows per session is not enough;
# we use a module-level cache keyed by (start_date, end_date).
_raw_cache = {'key': None, 'rows': [], 'columns': []}

# channel_rows / channel_month_rows are a pure function of the (already cached) raw SAP
# rows, but were re-aggregated on every /api/sales-data/ hit — two full O(n) passes over
# tens of thousands of invoice lines (plus a strptime per row in the month pass) before the
# response could be sent. Memoize them by date range so repeat loads / Fetch clicks within
# the SAP cache window skip the rework. Same shape/TTL as services._SALES_CACHE.
_CHANNEL_AGG_CACHE = {}        # 'start|end' -> (expires_at, channel_rows, channel_month_rows)
_CHANNEL_AGG_TTL = 90          # seconds

EDIT_PIN = 'gill'


def _get_type_filter(request):
    from core.context_processors import derive_realise_profile
    _, type_filter, _ = derive_realise_profile(request.user)
    return type_filter


def _parse_body(request):
    try:
        return json.loads(request.body)
    except Exception:
        return {}


@never_cache
@group_required(*REALISE_GROUPS, json_response=False)
def dashboard(request):
    # never_cache: the territory_payload (person map + per-channel state whitelist) is
    # baked into the HTML at render time, so the browser must re-fetch the page after a
    # mapping change instead of serving a stale copy (else newly-assigned states like a
    # freshly-added ECOM/NAGALAND wouldn't appear until a hard refresh).
    return render(request, 'realise/dashboard.html', {
        'sidebar_active': 'realise',
        'territory_payload': json.dumps(services.get_territory_dashboard_payload()),
    })


@permission_flag_required('can_oih_vs_stock')
def oih_vs_stock(request):
    """Standalone tab: open-order litres (OIH) vs warehouse stock per product, with the
    Required (OIH − Stock) gap. Reuses the /api/oih-breakdown/ data (OIH rows + per-item
    on-hand stock across the three warehouses)."""
    return render(request, 'realise/oih_vs_stock.html', {'sidebar_active': 'oih_vs_stock'})


@permission_flag_required('can_compare_sales')
def compare_sales(request):
    """Standalone tab: month-wise sales pivot (rows = chosen dimension, columns = months)
    with a Main Group filter (compare groups for the same period) and a Compare selector
    (Litres / Realise / Both). Reuses /api/sales-data/ channel_month_rows."""
    return render(request, 'realise/compare_sales.html', {
        'sidebar_active': 'compare_sales',
        'territory_payload': json.dumps(services.get_territory_dashboard_payload()),
    })


@permission_flag_required('can_sales_cn')
def sales_cn(request):
    """Standalone tab: gross Sales vs Credit Notes. Rows = a chosen dimension (main group /
    state / sales person / product / item / customer); columns = Total Sales, Total CN (split
    into CN for Goods and Claim for Services) and Net Sales = Total Sales − Total CN. Filters:
    company (Oil / Beverages), Revenue vs quantity (Litres/Boxes), Premium/Commodity (oil),
    and a date range. Data via /realise/api/sales-cn/. The territory payload lets the client
    resolve the Contact Person dimension to the mapped territory owner (same as Compare Sales)."""
    return render(request, 'realise/sales_cn.html', {
        'sidebar_active': 'sales_cn',
        'territory_payload': json.dumps(services.get_territory_dashboard_payload()),
    })


@permission_flag_required('can_customer_aging')
def customer_aging(request):
    """Standalone tab: customer-receivables aging pivot (FORMAT → customers) with the five
    aging buckets, computed live from SAP (B1 reconciliation logic) as of a selectable
    date (?as_of=YYYY-MM-DD, default today)."""
    from datetime import date, datetime
    today = date.today()
    try:
        aging_date = datetime.strptime(request.GET.get('as_of', ''), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        aging_date = today
    if aging_date > today:                  # no aging into the future
        aging_date = today
    remark_idx = services.get_aging_remark_index()
    return render(request, 'realise/customer_aging.html', {
        'sidebar_active': 'customer_aging',
        # raw dict — the template's |json_script does the JSON serialization (passing a
        # pre-dumped string here would double-encode and JSON.parse would yield a string).
        'aging_payload': services.get_customer_aging(aging_date),
        'aging_remark_index': remark_idx['index'],      # {card_code: [remark tokens]}
        'aging_remark_options': remark_idx['options'],  # sorted master list for the filter
        'aging_date': aging_date.isoformat(),
        'aging_today': today.isoformat(),
    })


@permission_flag_required('can_customer_aging')
def customer_aging_detail(request):
    """Full-page per-document detail behind one customer's Balance Due, with an editable
    Remarks column and a Pivot/Unpivot toggle. ?code=<CardCode>&name=<CardName>&as_of=YYYY-MM-DD."""
    from datetime import date, datetime
    today = date.today()
    try:
        aging_date = datetime.strptime(request.GET.get('as_of', ''), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        aging_date = today
    if aging_date > today:
        aging_date = today
    code = (request.GET.get('code') or '').strip()
    name = (request.GET.get('name') or '').strip() or code
    # Company: 'oil' (default) uses SAP_SCHEMA; 'mart' swaps to MART_SCHEMA and prefixes stored
    # remarks with 'MART:' so they never collide with oil's for the same CardCode/TransId.
    is_mart = (request.GET.get('company') or '').strip().lower() == 'mart'
    schema = services.MART_SCHEMA if is_mart else None
    rk_prefix = 'MART:' if is_mart else ''
    return render(request, 'realise/customer_aging_detail.html', {
        'sidebar_active': 'customer_aging',
        'detail_payload': {'code': code, 'name': name, 'aging_date': aging_date.isoformat(),
                           'company': 'mart' if is_mart else 'oil',
                           'categories': services.AGING_REMARK_CATEGORIES,
                           'grace_days': services.get_aging_grace_days(code) if code else 0,
                           'rows': services.get_customer_aging_detail(
                               code, aging_date, schema=schema, row_key_prefix=rk_prefix,
                               classify=not is_mart) if code else []},
        'aging_date': aging_date.isoformat(),
        'aging_today': today.isoformat(),
    })


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['GET'])
def api_customer_aging_beverages(request):
    """Raw open-invoice aging rows for the Jivo Beverages company, for the Beverages toggle on
    Customer Aging. The client pivots them (Sales Person → Customer) and offers the per-day
    multi-select + Excel-like raw drill. ?as_of=YYYY-MM-DD (default today)."""
    from datetime import date, datetime
    today = date.today()
    try:
        aging_date = datetime.strptime(request.GET.get('as_of', ''), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        aging_date = today
    if aging_date > today:
        aging_date = today
    return JsonResponse({'status': 'ok', **services.get_customer_aging_beverages(aging_date)})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['GET'])
def api_customer_aging_oil_ar(request):
    """Oil open-invoice RAW DATA rows (same shape as the Beverages endpoint) that back the oil
    RAW DATA workspace on Customer Aging. ?as_of=YYYY-MM-DD (default today)."""
    from datetime import date, datetime
    today = date.today()
    try:
        aging_date = datetime.strptime(request.GET.get('as_of', ''), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        aging_date = today
    if aging_date > today:
        aging_date = today
    return JsonResponse({'status': 'ok', **services.get_customer_aging_oil_ar(aging_date)})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['GET'])
def api_customer_aging_mart(request):
    """Raw open-invoice aging rows for the Jivo Mart company, for the Mart toggle on Customer
    Aging. Same shape/behaviour as the Beverages endpoint. ?as_of=YYYY-MM-DD (default today)."""
    from datetime import date, datetime
    today = date.today()
    try:
        aging_date = datetime.strptime(request.GET.get('as_of', ''), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        aging_date = today
    if aging_date > today:
        aging_date = today
    return JsonResponse({'status': 'ok', **services.get_customer_aging_mart(aging_date)})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark(request):
    """Save (or clear) one per-document remark on the Customer Aging detail page."""
    body = _parse_body(request)
    code = body.get('code', '')
    row_key = body.get('row_key', '')
    if not code or not row_key:
        return JsonResponse({'status': 'error', 'error': 'code and row_key required'}, status=400)
    services.save_aging_remark(code, row_key, body.get('remark', ''))
    return JsonResponse({'status': 'ok'})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_due_days(request):
    """Set (or clear) the NOT DUE grace period for one customer on the Customer Aging detail
    page. Body: {code:<CardCode>, days:<int>}. days<=0 turns the auto NOT DUE/OVERDUE off."""
    body = _parse_body(request)
    code = (body.get('code') or '').strip()
    if not code:
        return JsonResponse({'status': 'error', 'error': 'code required'}, status=400)
    days = services.save_aging_grace_days(code, body.get('days'), user=request.user)
    return JsonResponse({'status': 'ok', 'days': days})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark_lines(request):
    """Replace the split breakdown (TDS / RTV / Claim / …) behind one open document on the
    Customer Aging detail page. Body: {code, row_key, lines:[{category,amount,remark}, ...]}."""
    body = _parse_body(request)
    code = body.get('code', '')
    row_key = body.get('row_key', '')
    if not code or not row_key:
        return JsonResponse({'status': 'error', 'error': 'code and row_key required'}, status=400)
    lines = body.get('lines', [])
    if not isinstance(lines, list):
        return JsonResponse({'status': 'error', 'error': 'lines must be a list'}, status=400)
    services.save_aging_remark_lines(code, row_key, lines)
    return JsonResponse({'status': 'ok'})


def _norm_docno(v):
    """A cell value → a clean Doc No string (Excel often reads doc numbers as floats)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]
    return s


def _parse_amount(v):
    """A cell → float; tolerates ₹, thousands commas and blanks (else 0.0)."""
    s = str('' if v is None else v).replace('₹', '').replace(',', '').strip()
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _extract_doc_entries(table):
    """A sheet (list of rows) → {doc_no: {'remark': str|None, 'splits': [{category, amount,
    remark}, ...]}}, or None if it has no header row with a Doc-No-like column plus a Remark or a
    Category column. The Doc column matches any header containing 'doc' (not a date); 'remark' any
    'remark' header, 'category' any 'categ' header, 'amount' an exact 'amount'/'amt' (or a contains
    match that isn't a document-total column like Original/Balance/Total Amount).

    Each document's Remark/Category value becomes BOTH the note AND a split whose category is that
    value — so uploading a 'Remark / Category' column fills the note and the on-screen Category
    together. A split's amount is taken from an explicit Amount column when present, else left None
    = 'allocate the document's full Balance Due' (resolved in bulk_update_aging_remarks). Document-
    total columns (Original/Balance/Total Amount) are NOT used as the split amount, so an exported
    aging sheet re-uploaded with categories still allocates the balance, not the original.

    An optional 'Actual Sales Person' column (Beverages raw drill) is captured per doc as
    entry['actual_sp'] when present, so a remark upload can also re-assign the actual sales
    person; if the column is absent it is simply skipped (no actual_sp key set)."""
    doc_i = rem_i = cat_i = amt_i = asp_i = header_idx = None
    for idx, row in enumerate(table):
        cols = [str(c or '').strip().lower() for c in row]
        d = next((j for j, c in enumerate(cols) if 'doc' in c and 'date' not in c), None)
        if d is None:
            continue
        r = next((j for j, c in enumerate(cols) if 'remark' in c), None)
        cat = next((j for j, c in enumerate(cols) if 'categ' in c), None)
        # Optional 'Actual Sales Person' override column — matched on 'actual' + sales/sp/person.
        asp = next((j for j, c in enumerate(cols)
                    if 'actual' in c and ('sales' in c or 'sp' in c or 'person' in c)), None)
        # Amount: prefer an exact split-amount header; else a contains-match that is NOT a
        # document-total column (Original/Balance/Total/Gross Amount) — so an EXPORTED aging
        # sheet (which carries Original/Balance amounts) is never mistaken for a split sheet.
        amt = next((j for j, c in enumerate(cols)
                    if c in ('amount', 'amt', 'amount (₹)', 'amount(₹)', 'split amount', 'amount to allocate')), None)
        if amt is None:
            amt = next((j for j, c in enumerate(cols) if 'amount' in c
                        and not any(w in c for w in ('original', 'balance', 'total', 'gross'))), None)
        # A real header row = a Doc column plus a Remark or a Category column. An Amount column
        # ALONE does not make it a header (exported sheets carry Original/Balance amounts).
        if r is not None or cat is not None:
            doc_i, rem_i, cat_i, amt_i, asp_i, header_idx = d, r, cat, amt, asp, idx
            break
    if header_idx is None:
        return None
    # The note (AgingRemark) comes from the Remarks column, falling back to Category; the split
    # CATEGORY comes from the Category column, falling back to Remarks. So a single "Remark /
    # Category" column drives BOTH the note and the on-screen Category. With no explicit Amount
    # column a split's amount is left None = "allocate the document's full Balance Due".
    note_i = rem_i if rem_i is not None else cat_i
    cat_src = cat_i if cat_i is not None else rem_i

    def cell(row, i):
        return '' if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()

    out = {}
    for row in table[header_idx + 1:]:
        doc = _norm_docno(row[doc_i]) if doc_i < len(row) else ''
        if not doc:
            continue
        entry = out.setdefault(doc, {'remark': None, 'splits': []})
        note = cell(row, note_i)
        catv = cell(row, cat_src)
        asp = cell(row, asp_i)                            # optional Actual Sales Person override
        if note:
            entry['remark'] = note
        if asp:
            entry['actual_sp'] = asp
        if catv:                                          # the value → a split (its category label)
            amt = None if amt_i is None else _parse_amount(row[amt_i] if amt_i < len(row) else '')
            entry['splits'].append({'category': catv, 'amount': amt, 'remark': ''})
    return out


def _parse_remark_upload(uploaded):
    """Read an uploaded .xlsx/.csv into {doc_no: {'remark', 'splits'}} (see _extract_doc_entries).
    Scans EVERY sheet (a 'Data' sheet first) for the one carrying a Doc No column alongside
    Remarks and/or Category+Amount columns — so both our own export and hand-kept books (data on a
    later sheet next to a pivot/summary) work. Raises ValueError if no sheet matches."""
    raw = uploaded.read()
    name = (getattr(uploaded, 'name', '') or '').lower()
    tables = []
    if name.endswith('.csv'):
        text = raw.decode('utf-8-sig', errors='replace')
        tables = [list(csv.reader(io.StringIO(text)))]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        order = sorted(wb.sheetnames, key=lambda n: n.strip().lower() != 'data')   # 'Data' first
        tables = [[list(r) for r in wb[sn].iter_rows(values_only=True)] for sn in order]
    for table in tables:
        parsed = _extract_doc_entries(table)
        if parsed is not None:
            return parsed
    raise ValueError('Could not find a "Doc No" column with "Remarks" and/or "Category" + "Amount" columns')


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark_upload(request):
    """Bulk-update per-document Remarks from an uploaded .xlsx/.csv, matching on Doc No.
    multipart: file=<xlsx/csv>, code=<CardCode>, as_of=YYYY-MM-DD (the aging date the sheet
    was exported for). Updates every open line sharing a Doc No; blanks are left unchanged."""
    code = (request.POST.get('code') or '').strip()
    upload = request.FILES.get('file')
    if not code or not upload:
        return JsonResponse({'status': 'error', 'error': 'code and file are required'}, status=400)
    try:
        doc_remarks = _parse_remark_upload(upload)
    except Exception as exc:
        return JsonResponse({'status': 'error', 'error': str(exc)}, status=400)
    result = services.bulk_update_aging_remarks(code, _parse_as_of(request.POST.get('as_of')), doc_remarks)
    return JsonResponse({'status': 'ok', **result})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark_upload_beverages(request):
    """Bulk-update Beverages open-invoice Remarks from an uploaded .xlsx/.csv, matching on Doc No
    (no customer needed — the code for each Doc No is resolved from the beverages aging rows).
    multipart: file=<xlsx/csv>, as_of=YYYY-MM-DD."""
    upload = request.FILES.get('file')
    if not upload:
        return JsonResponse({'status': 'error', 'error': 'file is required'}, status=400)
    try:
        doc_remarks = _parse_remark_upload(upload)
    except Exception as exc:
        return JsonResponse({'status': 'error', 'error': str(exc)}, status=400)
    result = services.bulk_update_beverages_remarks(_parse_as_of(request.POST.get('as_of')), doc_remarks)
    status = 'error' if result.get('error') else 'ok'
    return JsonResponse({'status': status, **result})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark_upload_oil(request):
    """Bulk-update Oil RAW DATA open-invoice Remarks from an uploaded .xlsx/.csv, matched on Doc No
    (the code for each Doc No is resolved from the oil raw-invoice rows). multipart: file, as_of."""
    upload = request.FILES.get('file')
    if not upload:
        return JsonResponse({'status': 'error', 'error': 'file is required'}, status=400)
    try:
        doc_remarks = _parse_remark_upload(upload)
    except Exception as exc:
        return JsonResponse({'status': 'error', 'error': str(exc)}, status=400)
    result = services.bulk_update_oil_ar_remarks(_parse_as_of(request.POST.get('as_of')), doc_remarks)
    status = 'error' if result.get('error') else 'ok'
    return JsonResponse({'status': status, **result})


@permission_flag_required('can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_aging_remark_clear(request):
    """Clear all saved per-document Remarks for one customer (split breakdowns are kept).
    Body: {code:<CardCode>}."""
    body = _parse_body(request)
    code = (body.get('code') or '').strip()
    if not code:
        return JsonResponse({'status': 'error', 'error': 'code required'}, status=400)
    return JsonResponse({'status': 'ok', 'cleared': services.clear_aging_remarks(code)})


@any_permission_flag('can_realise', 'can_customer_aging', 'can_oih_vs_stock', 'can_compare_sales',
                     'can_claims', json_response=True)
@require_http_methods(['POST'])
def api_export_xlsx(request):
    """Build a multi-sheet .xlsx from client-supplied sheets and stream it back.
    Body: {filename, sheets:[{name, rows:[[cell, ...], ...]}, ...]} where each cell is a
    scalar (numbers become real numeric cells) or {value, style, colspan}. Generic — powers
    the Customer Aging detail 'Export Excel' (Pivot + Data sheets in one file)."""
    body = _parse_body(request)
    sheets_in = body.get('sheets') or []
    if not isinstance(sheets_in, list) or not sheets_in:
        return JsonResponse({'error': 'sheets required'}, status=400)
    sheets = []
    for s in sheets_in:
        if not isinstance(s, dict):
            continue
        rows = s.get('rows')
        if isinstance(rows, list) and rows:
            sheets.append((str(s.get('name') or 'Sheet'), rows))
    if not sheets:
        return JsonResponse({'error': 'no rows to export'}, status=400)
    content = build_workbook(sheets)
    filename = re.sub(r'[^A-Za-z0-9._ -]', '_', str(body.get('filename') or 'export'))[:120]
    if not filename.lower().endswith('.xlsx'):
        filename += '.xlsx'
    response = HttpResponse(
        content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _parse_as_of(s):
    """Parse a ?as_of=YYYY-MM-DD query param into a date, or None (→ today) if absent/bad."""
    try:
        return datetime.strptime((s or '').strip(), '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


@any_permission_flag('can_realise', 'can_customer_aging', json_response=True)
@require_http_methods(['POST'])
def api_export_aging_detail(request):
    """Whole-book Customer Aging detail export: every open document — with its per-document
    Remark and split breakdown — for the parties the client passes (the currently-filtered
    set), in one sheet. All detail is fetched in a single bulk SAP query."""
    body = _parse_body(request)
    parties = body.get('parties') or []
    if not isinstance(parties, list) or not parties:
        return JsonResponse({'error': 'parties required'}, status=400)
    codes, meta = [], {}
    for p in parties:
        if not isinstance(p, dict):
            continue
        code = str(p.get('code') or '').strip()
        if not code or code in meta:
            continue
        meta[code] = (str(p.get('name') or code).strip(), str(p.get('format') or '').strip())
        codes.append(code)
    if not codes:
        return JsonResponse({'error': 'no valid parties'}, status=400)

    as_of = _parse_as_of(body.get('as_of'))
    detail = services.get_customer_aging_detail_bulk(codes, as_of)

    def famt(a):
        try:
            return '{:,.0f}'.format(float(a or 0))
        except (TypeError, ValueError):
            return '0'

    hdr = {'bold': True, 'fill': '0F172A', 'color': 'FFFFFF'}
    headers = ['Format', 'Customer', 'Doc No', 'Type', 'Posting Date', 'Due Date', 'Branch',
               'Original', 'Balance Due', '0-30', '31-60', '61-90', '91-120', '121+', 'Remark', 'Splits']
    rows = [[{'value': h, **hdr} for h in headers]]
    tot = {k: 0.0 for k in ('original', 'balance_due', 'b0_30', 'b31_60', 'b61_90', 'b91_120', 'b121')}
    for code in codes:
        name, fmt = meta[code]
        for d in detail.get(code, []):
            splits = '; '.join(
                (('%s: %s' % (s.get('category') or '?', famt(s.get('amount'))))
                 + ((' (%s)' % s['remark']) if s.get('remark') else ''))
                for s in (d.get('splits') or []))
            rows.append([
                fmt, name, d['doc_no'], d['type'], d['posting_date'], d['due_date'], d['branch'],
                d['original'], d['balance_due'], d['b0_30'], d['b31_60'], d['b61_90'],
                d['b91_120'], d['b121'], d['remark'], splits,
            ])
            for k in tot:
                tot[k] += d.get(k, 0) or 0
    if len(rows) == 1:
        return JsonResponse({'error': 'No open documents for the selected parties'}, status=400)

    def tcell(v):
        return {'value': round(v, 2), 'bold': True, 'fill': 'E2E8F0'}
    rows.append([{'value': 'TOTAL', 'bold': True, 'fill': 'E2E8F0'}] + ['' ] * 6
                + [tcell(tot['original']), tcell(tot['balance_due']), tcell(tot['b0_30']),
                   tcell(tot['b31_60']), tcell(tot['b61_90']), tcell(tot['b91_120']), tcell(tot['b121'])]
                + ['', ''])

    content = build_workbook([('Aging Detail', rows)])
    fname = 'Customer Aging Detail %s.xlsx' % (as_of or datetime.now().date()).strftime('%d.%m.%Y')
    resp = HttpResponse(
        content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="%s"' % fname
    return resp


@permission_flag_required('can_required_credit_limit')
def required_credit_limit(request):
    """Standalone tab: Required Credit Limit — live Order-in-Hand grouped by ASM
    (territory owner) → party, in the closing-sheet layout. Each party row shows open
    litres + open value (₹), a Premium/Commodity filter at the top, and a frontend-
    editable delivery remark (the only writable column; persisted to ClosingRemark).
    ?as_of=YYYY-MM-DD (default today) drives the Ledger Amt / Payment Done date view."""
    return render(request, 'realise/required_credit_limit.html', {
        'sidebar_active': 'required_credit_limit',
        # raw dict — the template's |json_script does the JSON serialization.
        'credit_payload': services.get_required_credit_rows(as_of_date=_parse_as_of(request.GET.get('as_of'))),
    })


@permission_flag_required('can_required_credit_limit', json_response=True)
@require_http_methods(['POST'])
def api_save_closing_remark(request):
    """Save the editable delivery remark for one party on the Required Credit Limit tab."""
    body = _parse_body(request)
    card_code = body.get('card_code', '')
    if not card_code:
        return JsonResponse({'status': 'error', 'error': 'card_code required'}, status=400)
    services.save_closing_remark(card_code, body.get('remark', ''), request.user)
    return JsonResponse({'status': 'ok'})


@permission_flag_required('can_required_credit_limit', json_response=True)
@require_http_methods(['POST'])
def api_credit_lock(request):
    """Freeze Total Outstanding + Required Limit at their current values for the chosen
    number of days. Snapshots every party row; returns the new lock state."""
    body = _parse_body(request)
    lock = services.create_credit_lock(body.get('days', 30), request.user)
    return JsonResponse({'status': 'ok', 'lock': lock})


@permission_flag_required('can_required_credit_limit', json_response=True)
@require_http_methods(['POST'])
def api_credit_unlock(request):
    """Lift the active lock early — the columns revert to live SAP immediately."""
    services.clear_credit_lock()
    return JsonResponse({'status': 'ok', 'lock': None})


@permission_flag_required('can_required_credit_limit')
@require_http_methods(['GET'])
def export_required_credit(request):
    """Download the Required Credit Limit data as an .xlsx in the CLOSING SHEET layout.
    ?type=P|C|P+C (repeatable) scopes to the on-screen Type filter; ?asm=<name> to one ASM."""
    type_filters = [t.strip() for t in request.GET.getlist('type') if t.strip() in ('P', 'C', 'P+C')]
    asms = [a.strip() for a in request.GET.getlist('asm') if a.strip()]
    payload = services.get_required_credit_rows(as_of_date=_parse_as_of(request.GET.get('as_of')))
    if asms:                     # scope the export to the on-screen ASM selection
        chosen = set(asms)
        payload = {'asms': [g for g in payload.get('asms', []) if g.get('asm') in chosen],
                   'total': payload.get('total', {})}
    content = services.build_closing_sheet_xlsx(payload, type_filters)
    parts = []
    if asms:
        parts.append(asms[0] if len(asms) == 1 else f'{len(asms)} ASMs')
    type_names = {'P': 'Premium', 'C': 'Commodity', 'P+C': 'Prem+Comm'}
    parts.append('+'.join(type_names[t] for t in type_filters) if type_filters else 'All')
    response = HttpResponse(
        content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="CLOSING SHEET ({" - ".join(parts)}).xlsx"'
    return response


def _aggregate_channel_rows(raw_rows):
    """Collapse raw SAP transaction rows to distinct
    (type, sub_group, main_group, state, sales_person, card_name, item_name) buckets
    with summed litres/revenue. Slide 2 only ever SUMS these dimensions, so this is
    lossless for every card / drill / commodity aggregation while still collapsing
    the many invoice LINES per customer-order into one bucket. card_name powers the
    Customer drill and item_name the Item Name drill in the channel detail modal."""
    agg = {}
    for row in raw_rows:
        sales_person = ''
        for k in ('U_SALES_PERSON', 'U_Sales_Person', 'SALES_PERSON', 'SalesPerson', 'SlpName'):
            v = str(row.get(k, '') or '').strip().upper()
            if v:
                sales_person = v
                break
        u_type = str(row.get('U_TYPE', '') or '').strip().upper()
        u_sub = str(row.get('U_Sub_Group', '') or '').strip().upper()
        u_main = str(row.get('U_Main_Group', '') or '').strip().upper()
        state = str(row.get('State', '') or '').strip().upper()
        card_name = str(row.get('CardName', '') or '').strip().upper()
        item_name = services._item_label(str(row.get('ItemCode', '') or '').strip().upper(),
                                          str(row.get('ItemName', '') or '').strip().upper())
        # SKU = OITM.U_SKU (pack size, e.g. '1 LTR' / '500 MLS'), returned by the proc as
        # "SKU". Carried per row so the Item-first drill can filter by SKU / Product.
        sku = str(row.get('SKU', '') or '').strip().upper()
        key = (u_type, u_sub, u_main, state, sales_person, card_name, item_name, sku)
        bucket = agg.get(key)
        if bucket is None:
            bucket = agg[key] = {
                'u_type': u_type, 'u_sub_group': u_sub, 'u_main_group': u_main,
                'state': state, 'sales_person': sales_person, 'card_name': card_name,
                'item_name': item_name, 'sku': sku,
                'liter': 0.0, 'line_total': 0.0,
            }
        bucket['liter'] += float(row.get('Liter', 0) or 0)
        bucket['line_total'] += float(row.get('LineTotal', 0) or 0)
    out = list(agg.values())
    for b in out:
        b['liter'] = round(b['liter'], 2)
        b['line_total'] = round(b['line_total'], 2)
    return out


def _aggregate_channel_month_rows(raw_rows):
    """Month-level channel buckets for the OILS month-wise pivot: distinct
    (type, main_group, state, sales_person, sub_group, item_name, card_name, ym) with
    summed litres. Carries every dimension the pivot's Drill By offers (State / Contact
    Person / Product / Item Name / Customer) plus the month, so the States x Months pivot
    can re-pivot its rows by any of them while months stay in the columns. `ym` is a
    sortable 'YYYY-MM'; `mlabel` is the display label ('JUL 2025')."""
    agg = {}
    for row in raw_rows:
        mon, year = services._parse_doc_date(row.get('DocDate', ''))
        if not mon or not year:
            continue
        try:
            mnum = datetime.strptime(mon, '%b').month
        except ValueError:
            continue
        sales_person = ''
        for k in ('U_SALES_PERSON', 'U_Sales_Person', 'SALES_PERSON', 'SalesPerson', 'SlpName'):
            v = str(row.get(k, '') or '').strip().upper()
            if v:
                sales_person = v
                break
        u_type = str(row.get('U_TYPE', '') or '').strip().upper()
        u_main = str(row.get('U_Main_Group', '') or '').strip().upper()
        u_sub = str(row.get('U_Sub_Group', '') or '').strip().upper()
        state = str(row.get('State', '') or '').strip().upper()
        item_name = services._item_label(str(row.get('ItemCode', '') or '').strip().upper(),
                                          str(row.get('ItemName', '') or '').strip().upper())
        card_name = str(row.get('CardName', '') or '').strip().upper()
        sku = str(row.get('SKU', '') or '').strip().upper()   # OITM.U_SKU pack size
        ym = '%s-%02d' % (year, mnum)
        key = (u_type, u_main, state, sales_person, u_sub, item_name, card_name, ym, sku)
        bucket = agg.get(key)
        if bucket is None:
            bucket = agg[key] = {
                'u_type': u_type, 'main_group': u_main, 'state': state,
                'sales_person': sales_person, 'u_sub_group': u_sub, 'item_name': item_name,
                'card_name': card_name, 'sku': sku, 'ym': ym, 'mlabel': '%s %s' % (mon, year),
                'liter': 0.0, 'line_total': 0.0,
            }
        bucket['liter'] += float(row.get('Liter', 0) or 0)
        bucket['line_total'] += float(row.get('LineTotal', 0) or 0)
    out = list(agg.values())
    for b in out:
        b['liter'] = round(b['liter'], 2)
        b['line_total'] = round(b['line_total'], 2)
    return out


def _channel_aggregates(start_date, end_date, raw_rows):
    """Memoized (channel_rows, channel_month_rows) for a date range. Only recomputes when
    the SAP cache window has rolled over; otherwise returns the prior aggregation so a
    cached sales-data load doesn't re-walk the full raw set on every request."""
    key = f'{start_date}|{end_date}'
    now = time.time()
    hit = _CHANNEL_AGG_CACHE.get(key)
    if hit and hit[0] > now:
        return hit[1], hit[2]
    channel_rows = _aggregate_channel_rows(raw_rows)
    channel_month_rows = _aggregate_channel_month_rows(raw_rows)
    if raw_rows:                          # only cache successful, non-empty pulls
        _CHANNEL_AGG_CACHE[key] = (now + _CHANNEL_AGG_TTL, channel_rows, channel_month_rows)
        for k in [k for k, v in _CHANNEL_AGG_CACHE.items() if v[0] <= now]:
            _CHANNEL_AGG_CACHE.pop(k, None)
    return channel_rows, channel_month_rows


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_health(request):
    ok, message = sap_connector.health_check()
    from core.context_processors import derive_realise_profile
    role, _, _ = derive_realise_profile(request.user)
    return JsonResponse({
        'sap_connected': ok,
        'message': message,
        'username': request.user.username,
        'role': role,
    })


@any_permission_flag('can_realise', 'can_compare_sales', json_response=True)
@require_http_methods(['POST'])
def api_sales_data(request):
    body = _parse_body(request)
    start_date = body.get('start_date', '')
    end_date   = body.get('end_date', '')
    if not start_date or not end_date:
        return JsonResponse({'status': 'error', 'error': 'start_date and end_date required'}, status=400)

    type_filter = _get_type_filter(request)
    force = bool(body.get('refresh') or body.get('force'))   # "Refresh from SAP" → fresh pull

    try:
        result, raw_rows = services.get_sales_data_cached(start_date, end_date, force=force)
    except Exception as e:
        logger.error('[REALISE] get_sales_data error: %s', e)
        return JsonResponse({'status': 'ok', 'data': [], 'count': 0})

    # Cache raw rows for drill-down and CSV export
    _raw_cache['key']     = f'{start_date}_{end_date}'
    _raw_cache['rows']    = raw_rows
    _raw_cache['columns'] = list(raw_rows[0].keys()) if raw_rows else []
    _raw_cache['start']   = start_date
    _raw_cache['end']     = end_date

    # Attach targets to each product row
    rows = result['products']
    month_year_pairs = set()
    for r in rows:
        if r.get('month') and r.get('year'):
            try:
                m = datetime.strptime(r['month'], '%b').month
                y = int(r['year'])
                month_year_pairs.add((m, y))
            except (ValueError, KeyError):
                pass

    # Build targets lookup for all month/year pairs found
    targets_cache = {}
    for (m, y) in month_year_pairs:
        targets_cache[(m, y)] = services.get_targets_for_month(m, y)

    output = []
    for r in rows:
        if type_filter and r['u_type'] != type_filter:
            continue

        tgt_ltrs = 0
        tgt_rate  = 0
        if r.get('month') and r.get('year'):
            try:
                m = datetime.strptime(r['month'], '%b').month
                y = int(r['year'])
                key = f"{r['u_type']}|{r['u_sub_group']}"
                td = targets_cache.get((m, y), {}).get(key, {})
                tgt_ltrs = td.get('tgt_ltrs', 0)
                tgt_rate  = td.get('tgt_rate', 0)
            except (ValueError, KeyError):
                pass

        output.append({
            'u_type':         r['u_type'],
            'u_sub_group':    r['u_sub_group'],
            'month':          r['month'],
            'year':           r['year'],
            'litres':         r['litres'],
            'linetotal':      r['linetotal'],
            'realise':        r['realise'],
            'target_sale':    tgt_ltrs,
            'target_realise': tgt_rate,
        })

    output.sort(key=lambda x: (
        0 if x['u_type'] == 'PREMIUM' else 1,
        -x.get('target_sale', 0),
        x['u_sub_group'],
        x['month'],
    ))

    channel_rows, channel_month_rows = _channel_aggregates(start_date, end_date, raw_rows)
    return JsonResponse({'status': 'ok', 'data': output, 'count': len(output),
                         'channel_rows': channel_rows, 'channel_month_rows': channel_month_rows})


@any_permission_flag('can_sales_cn', json_response=True)
@require_http_methods(['POST'])
def api_sales_cn_data(request):
    """Sales vs Credit Notes rows for a date range + company. Body: {start_date, end_date,
    company: 'oil'|'beverages'}. Returns the get_sales_cn_report payload (rows carry each
    dimension plus the sales / CN-goods / CN-service measures; the client pivots + filters)."""
    body = _parse_body(request)
    start_date = body.get('start_date', '')
    end_date = body.get('end_date', '')
    company = body.get('company', 'oil')
    if not start_date or not end_date:
        return JsonResponse({'status': 'error', 'error': 'start_date and end_date required'}, status=400)
    return JsonResponse(services.get_sales_cn_report(start_date, end_date, company))


@permission_flag_required('can_hidden_sales')
def hidden_sales(request):
    """Standalone tab: sales invoices flagged HIDDEN (OINV.U_ARNO='H') — the ones excluded from
    the dashboard's Done — surfaced per invoice line and drillable by customer / item / cost
    center / date / status, with quantity, litres and value. Data via /realise/api/hidden-sales/."""
    return render(request, 'realise/hidden_sales.html', {'sidebar_active': 'hidden_sales'})


@any_permission_flag('can_hidden_sales', json_response=True)
@require_http_methods(['POST'])
def api_hidden_sales_data(request):
    """Hidden invoice lines for a date range. Body: {start_date, end_date}."""
    body = _parse_body(request)
    start_date = body.get('start_date', '')
    end_date = body.get('end_date', '')
    if not start_date or not end_date:
        return JsonResponse({'status': 'error', 'error': 'start_date and end_date required'}, status=400)
    return JsonResponse(services.get_hidden_customer_sales(start_date, end_date))


@permission_flag_required('can_customer_master')
def customer_master(request):
    """Standalone tab: the customer master — every customer (OCRD) with contact details, GSTIN /
    PAN, address & location, sales person, payment terms, credit limit, balance and status.
    Searchable / filterable table with an Excel export. Data via /realise/api/customer-master/."""
    return render(request, 'realise/customer_master.html', {'sidebar_active': 'customer_master'})


@any_permission_flag('can_customer_master', json_response=True)
@require_http_methods(['GET'])
def api_customer_master_data(request):
    """Full customer master (all OCRD customers) as JSON. Cached in the service layer."""
    return JsonResponse(services.get_customer_master())


# (key, column header) for the Customer Master Excel export — order = on-screen order.
_CUST_MASTER_COLS = [
    ('code', 'Code'), ('name', 'Customer Name'), ('main_group', 'Main Group'),
    ('status', 'Status'), ('gstin', 'GSTIN'), ('pan', 'PAN'),
    ('contact_person', 'Contact Person'), ('mobile', 'Mobile'),
    ('email', 'Email'), ('address', 'Address'), ('city', 'City'), ('state', 'State'),
    ('pincode', 'Pincode'), ('sales_person', 'Sales Person'), ('payment_terms', 'Payment Terms'),
    ('credit_limit', 'Credit Limit'), ('balance', 'Balance'),
]


@permission_flag_required('can_customer_master')
@require_http_methods(['GET'])
def export_customer_master(request):
    """Download the customer master as an .xlsx (all columns, all customers)."""
    payload = services.get_customer_master()

    def _cell(key, r):
        v = r.get(key, '')
        return round(float(v or 0), 2) if key in ('credit_limit', 'balance') else v

    header = [label for _, label in _CUST_MASTER_COLS]
    body = [[_cell(key, r) for key, _ in _CUST_MASTER_COLS] for r in payload.get('rows', [])]
    content = build_workbook([('Customer Master', [header] + body)])
    response = HttpResponse(
        content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Customer Master.xlsx"'
    return response


@permission_flag_required('can_claims')
def claims(request):
    """Standalone tab: the Claims register — a manually-maintained table of claims (Claim Date,
    Party, Pass Date, Month & Year, Type, Hold, Amount, and the manual Passed / Hold / Reason
    columns). Rows are added/edited in-app (nothing is read from SAP as report data); SAP only
    powers the party and product/item pickers. Filter by Month/Year, date range and Type; Drill By
    Customer / Product / Item / Main Group. Data via /realise/api/claims/."""
    return render(request, 'realise/claims.html', {'sidebar_active': 'claims'})


@any_permission_flag('can_claims', json_response=True)
@require_http_methods(['GET'])
def api_claims_data(request):
    """All claim rows plus the entry-picker masters (customers / products / items)."""
    payload = services.get_claims()
    payload['masters'] = services.get_claim_masters()
    return JsonResponse(payload)


@any_permission_flag('can_claims', json_response=True)
@require_http_methods(['POST'])
def api_claim_save(request):
    """Create or update one claim from the add/edit form. Body: the claim fields (+ optional id)."""
    body = _parse_body(request)
    try:
        row = services.upsert_claim(body, user=request.user)
    except ValueError as exc:
        return JsonResponse({'status': 'error', 'error': str(exc)}, status=400)
    except Exception as exc:
        logger.error('[CLAIMS] save failed: %s', exc)
        return JsonResponse({'status': 'error', 'error': 'Could not save the claim.'}, status=500)
    return JsonResponse({'status': 'ok', 'row': row})


@any_permission_flag('can_claims', json_response=True)
@require_http_methods(['POST'])
def api_claim_delete(request):
    """Delete one claim by id. Body: {id}."""
    body = _parse_body(request)
    cid = body.get('id')
    if not cid:
        return JsonResponse({'status': 'error', 'error': 'id required'}, status=400)
    return JsonResponse({'status': 'ok', 'deleted': services.delete_claim(cid)})


@permission_flag_required('can_sales_flow')
def sales_document_flow(request):
    """Standalone tab: the sales document chain for a day's sales — Party, Sales Quotation No,
    Sales Order No, Invoice No and invoiced Litres, one row per document chain. Defaults to
    yesterday. Data via /realise/api/sales-flow/."""
    return render(request, 'realise/sales_document_flow.html', {'sidebar_active': 'sales_flow'})


@any_permission_flag('can_sales_flow', json_response=True)
@require_http_methods(['POST'])
def api_sales_flow_data(request):
    """Sales document-flow rows for a date range (defaults to yesterday when omitted).
    Body: {start_date, end_date}."""
    from datetime import date, timedelta
    body = _parse_body(request)
    yday = (date.today() - timedelta(days=1)).isoformat()
    start_date = body.get('start_date') or yday
    end_date = body.get('end_date') or yday
    company = body.get('company', 'oil')
    return JsonResponse(services.get_sales_document_flow(start_date, end_date, company))


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['POST'])
def api_beverages_data(request):
    """Beverages dataset (JIVO_BEVERAGES_HANADB / REPORT_SALES_COGS) — granular rows
    by Variety / Sub-Group / SKU with Quantity & Boxes for the dynamic driller."""
    body = _parse_body(request)
    start_date = body.get('start_date', '')
    end_date   = body.get('end_date', '')
    if not start_date or not end_date:
        return JsonResponse({'status': 'error', 'error': 'start_date and end_date required'}, status=400)
    try:
        data = services.get_beverages_rows_cached(start_date, end_date)
    except Exception as e:
        logger.error('[BEVERAGES] fetch error: %s', e)
        return JsonResponse({'status': 'ok', 'data': [], 'count': 0, 'today_boxes': 0, 'yesterday_boxes': 0,
                             'today_items': [], 'yesterday_items': [], 'today_date': '', 'yesterday_date': '',
                             'customer_rows': [], 'month_rows': [], 'oih_rows': []})
    is_dict = isinstance(data, dict)
    rows = data.get('rows', []) if is_dict else (data or [])
    return JsonResponse({'status': 'ok', 'data': rows, 'count': len(rows),
                         'today_boxes': data.get('today_boxes', 0) if is_dict else 0,
                         'yesterday_boxes': data.get('yesterday_boxes', 0) if is_dict else 0,
                         'today_items': data.get('today_items', []) if is_dict else [],
                         'yesterday_items': data.get('yesterday_items', []) if is_dict else [],
                         'today_date': data.get('today_date', '') if is_dict else '',
                         'yesterday_date': data.get('yesterday_date', '') if is_dict else '',
                         'customer_rows': data.get('customer_rows', []) if is_dict else [],
                         'month_rows': data.get('month_rows', []) if is_dict else [],
                         'oih_rows': data.get('oih_rows', []) if is_dict else []})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_beverages_docs(request):
    """Invoice / open-SO documents behind a beverages driller cell, filtered to the clicked
    node (customer + ancestor dims + brand/month). metric=sales -> invoices, oih -> SOs."""
    start = request.GET.get('start', '') or ''
    end = request.GET.get('end', '') or ''
    metric = str(request.GET.get('metric', 'sales') or 'sales').strip().lower()
    if not start or not end:
        return JsonResponse({'status': 'error', 'error': 'start and end required'}, status=400)
    filters = {}
    for key in ('variety', 'sub_group', 'sku', 'item', 'main_group', 'state',
                'brand', 'chain', 'sales_person', 'customer', 'ym'):
        val = request.GET.get('f_' + key)
        if val not in (None, ''):
            filters[key] = str(val).strip().upper()
    try:
        data = services.get_beverages_documents(start, end, filters, metric)
    except Exception as exc:
        logger.error('[BEVERAGES] docs fetch failed: %s', exc)
        return JsonResponse({'status': 'ok', 'metric': metric, 'count': 0, 'data': []})
    return JsonResponse({'status': 'ok', 'metric': metric, 'count': len(data), 'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['POST'])
def api_drill_down(request):
    body = _parse_body(request)
    start_date   = body.get('start_date', '')
    end_date     = body.get('end_date', '')
    u_type       = body.get('u_type') or body.get('product_type', '')
    u_sub_group  = body.get('u_sub_group') or body.get('sub_group', '')
    drill_by     = body.get('drill_by', 'State')
    month        = body.get('month')
    year         = body.get('year')
    filters      = body.get('filters') or {}

    type_filter = _get_type_filter(request)
    if type_filter and u_type and u_type.upper() != type_filter:
        return JsonResponse({'data': []})

    cache_key = f'{start_date}_{end_date}'
    if _raw_cache['key'] != cache_key:
        try:
            _, raw_rows = services.get_sales_data(start_date, end_date)
            _raw_cache['key']  = cache_key
            _raw_cache['rows'] = raw_rows
        except Exception as e:
            logger.error('[DRILL] fetch failed: %s', e)
            return JsonResponse({'data': []})
    else:
        raw_rows = _raw_cache['rows']

    data = services.get_drill_down(
        start_date, end_date, raw_rows,
        u_type=u_type, u_sub_group=u_sub_group,
        drill_by=drill_by, month=month, year=year, filters=filters,
    )
    return JsonResponse({'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['POST'])
def api_historical_realise(request):
    body = _parse_body(request)
    start_date = body.get('start_date', '')
    end_date   = body.get('end_date', '')
    period     = body.get('period', '12m')

    type_filter = _get_type_filter(request)

    try:
        result, drill_result = services.get_historical_realise(start_date, end_date, period)
    except Exception as e:
        logger.error('[HIST] error: %s', e)
        return JsonResponse({'status': 'ok', 'data': {}, 'drill_data': {}})

    if type_filter:
        result = {k: v for k, v in result.items() if k.startswith(type_filter + '|')}
        drill_result = {k: v for k, v in drill_result.items() if k.startswith(type_filter + '|')}

    return JsonResponse({'status': 'ok', 'data': result, 'drill_data': drill_result, 'period': period})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_targets(request):
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year  = int(request.GET.get('year',  datetime.now().year))
    except (ValueError, TypeError):
        month = datetime.now().month
        year  = datetime.now().year

    data = services.get_targets_for_month(month, year)
    return JsonResponse({'status': 'ok', 'month': month, 'year': year, 'data': data})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_save_targets(request):
    body = _parse_body(request)
    month   = body.get('month')
    year    = body.get('year')
    targets = body.get('targets', [])

    if not month or not year or not targets:
        return JsonResponse({'status': 'error', 'error': 'month, year, and targets required'}, status=400)

    try:
        month = int(month)
        year  = int(year)
    except (ValueError, TypeError):
        return JsonResponse({'status': 'error', 'error': 'month and year must be integers'}, status=400)

    saved = services.save_monthly_targets(targets, month, year, request.user)
    return JsonResponse({'status': 'ok', 'saved': saved})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['POST'])
def api_verify_pin(request):
    body = _parse_body(request)
    pin = body.get('pin', '')
    if pin == EDIT_PIN:
        return JsonResponse({'status': 'ok', 'verified': True})
    return JsonResponse({'status': 'error', 'verified': False, 'message': 'Incorrect password'})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_channel_targets(request):
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month = datetime.now().month
        year = datetime.now().year
    data = services.get_channel_target_map(month, year, request.GET.get('seg', ''))
    return JsonResponse({'status': 'ok', 'month': month, 'year': year, 'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_order_in_hand(request):
    data = services.get_order_in_hand_by_person()
    return JsonResponse({'status': 'ok', 'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_order_in_hand_rows(request):
    return JsonResponse({'status': 'ok', 'data': services.get_order_in_hand_rows()})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_sales_pulse(request):
    """Tiny data fingerprint for the dashboard 'live' heartbeat (see services.get_sales_pulse).
    The client polls this cheaply every ~30s and only forces a fresh pull when it changes."""
    dataset = (request.GET.get('dataset') or 'oils').strip().lower()
    start = (request.GET.get('start') or '').strip()
    end = (request.GET.get('end') or '').strip()
    if not start or not end:
        return JsonResponse({'status': 'ok', 'pulse': ''})
    return JsonResponse({'status': 'ok', 'pulse': services.get_sales_pulse_cached(dataset, start, end)})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_channel_detail_docs(request):
    """Documents (invoices / open SOs) behind a Done-L or Order-in-Hand cell in the
    channel-detail modal, filtered to the clicked drill node. Same SAP sources as the
    dashboard's Done / Order-in-Hand numbers."""
    channel = str(request.GET.get('channel', '') or '').strip().upper()
    metric = str(request.GET.get('metric', 'done') or 'done').strip().lower()
    seg = request.GET.get('seg', '')
    filters = {}
    for key in ('group', 'state', 'person', 'customer', 'product', 'item'):
        val = request.GET.get(key)
        if val not in (None, ''):
            filters[key] = str(val).strip().upper()
    if metric == 'oih':
        data = services.get_channel_oih_documents(channel, filters, seg)
    else:
        start = request.GET.get('start') or _raw_cache.get('start') or ''
        end = request.GET.get('end') or _raw_cache.get('end') or ''
        # Diagnostic: ?reconcile=1 returns a party-by-party comparison of the channel Done
        # (proc) vs the popup Done (direct query) so any mismatch can be pinpointed.
        if request.GET.get('reconcile'):
            return JsonResponse({'status': 'ok', 'reconcile':
                services.reconcile_channel_done(start, end, channel, seg, filters.get('state', ''))})
        data = services.get_channel_done_documents(start, end, channel, seg, filters)
    return JsonResponse({'status': 'ok', 'metric': metric, 'count': len(data),
                         'warehouses': services.OIH_STOCK_WAREHOUSES, 'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_commodity_oih_rows(request):
    """Open-order litres for COMMODITY items, shaped for the commodity table's OIH."""
    return JsonResponse({'status': 'ok', 'data': services.get_commodity_oih_rows()})


@any_permission_flag('can_realise', 'can_oih_vs_stock', json_response=True)
@require_http_methods(['GET'])
def api_oih_breakdown(request):
    """Granular open-order litres by item dimensions (split Premium/Commodity) for the
    OIH KPI window's dynamic drill; the client nests them into any chosen order.
    Cached (90s) so repeat opens of the OIH-vs-Stock tab / dashboard reuse one pull."""
    return JsonResponse({'status': 'ok', **services.get_oih_dimension_rows_cached()})


@any_permission_flag('can_realise', 'can_oih_vs_stock', json_response=True)
@require_http_methods(['GET'])
def api_oih_breakdown_beverages(request):
    """Jivo Beverages variant of the OIH-vs-Stock breakdown: open-order BOXES vs on-hand
    stock (boxes), grouped by the beverages family (DRINKS/WATER/…) and variety. Same payload
    shape as api_oih_breakdown so the OIH-vs-Stock page renders it with the Beverages toggle."""
    return JsonResponse({'status': 'ok', **services.get_oih_dimension_rows_beverages_cached()})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_target_nodes(request):
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    return JsonResponse({'status': 'ok', 'month': month, 'year': year,
                         'data': services.get_target_nodes(month, year, request.GET.get('seg', ''))})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET', 'POST'])
def api_flex_targets(request):
    """Persisted 'Flex TGT' overrides for the Sales Channel drill table.
    GET  ?seg=&month=&year=          -> {data: {row_key: value}}
    POST {seg, month, year, row_key, value}  upserts one override (value null/'' clears it)."""
    if request.method == 'POST':
        body = _parse_body(request)
        ok = services.save_flex_target(body.get('seg', ''), body.get('month'), body.get('year'),
                                       body.get('row_key', ''), body.get('value'))
        return JsonResponse({'status': 'ok' if ok else 'error'})
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    return JsonResponse({'status': 'ok', 'month': month, 'year': year,
                         'data': services.get_flex_targets(request.GET.get('seg', ''), month, year)})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_segment_targets(request):
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month = datetime.now().month
        year = datetime.now().year

    segment = str(request.GET.get('segment', 'state') or 'state').strip().lower()
    data = services.get_segment_target_map(segment, month, year, request.GET.get('seg', ''))
    return JsonResponse({'status': 'ok', 'segment': segment, 'month': month, 'year': year, 'data': data})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_territory_map(request):
    """The fixed channel×state grid with each cell's current sales person, plus the
    channel order and the distinct people list (for the reassign dropdown)."""
    return JsonResponse({'status': 'ok', **services.get_territory_map_payload()})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_save_territory_map(request):
    """Update ONLY the sales_person on existing grid cells (admin only)."""
    body = _parse_body(request)
    assignments = body.get('assignments', [])
    if not isinstance(assignments, list):
        return JsonResponse({'status': 'error', 'error': 'assignments must be a list'}, status=400)
    saved = services.save_territory_persons(assignments, request.user)
    return JsonResponse({'status': 'ok', 'saved': saved})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_refresh_territory_map(request):
    """Re-sync the fixed grid from live SAP (adds new channel×state cells; never
    wipes existing person assignments). Admin only."""
    from django.core.management import call_command
    try:
        call_command('seed_territory_map', refresh=True)
    except Exception as e:
        logger.error('[TERRITORY] refresh failed: %s', e)
        return JsonResponse({'status': 'error', 'error': str(e)}, status=500)
    return JsonResponse({'status': 'ok', **services.get_territory_map_payload()})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_territory_targets(request):
    """Single target (litres) per (channel, state) territory for a month/year."""
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    return JsonResponse({'status': 'ok', 'month': month, 'year': year,
                         'data': services.get_territory_targets(month, year)})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_save_territory_targets(request):
    """Save one target (litres) per (channel, state) territory (admin only)."""
    body = _parse_body(request)
    month, year = body.get('month'), body.get('year')
    targets = body.get('targets', [])
    if not month or not year or not isinstance(targets, list):
        return JsonResponse({'status': 'error', 'error': 'month, year, targets[] required'}, status=400)
    try:
        month, year = int(month), int(year)
    except (ValueError, TypeError):
        return JsonResponse({'status': 'error', 'error': 'month and year must be integers'}, status=400)
    saved = services.save_territory_targets(month, year, targets)
    return JsonResponse({'status': 'ok', 'saved': saved})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_territory_product_targets(request):
    """Per-product targets (litres+realise) per (channel,state) for a period, plus the
    product master — the shape the Person Mapping 'Set product targets' UI consumes."""
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    return JsonResponse({'status': 'ok', 'month': month, 'year': year,
                         'products': services.get_product_master(),
                         'data': services.get_territory_product_targets(month, year)})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_save_territory_product_targets(request):
    """Save the full per-product target set for a period; rolls up into the dashboard's
    TargetNode (channel/state) and MonthlyTarget (per-product). Admin only."""
    body = _parse_body(request)
    month, year = body.get('month'), body.get('year')
    targets = body.get('targets', {})
    if not month or not year or not isinstance(targets, dict):
        return JsonResponse({'status': 'error', 'error': 'month, year, targets{} required'}, status=400)
    try:
        month, year = int(month), int(year)
    except (ValueError, TypeError):
        return JsonResponse({'status': 'error', 'error': 'month and year must be integers'}, status=400)
    saved = services.save_territory_product_targets(month, year, targets, request.user)
    return JsonResponse({'status': 'ok', 'saved': saved})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_channel_quick_targets(request):
    """Channel Targets editor data: each channel's previous-month actual sale +
    current channel-level target for the selected (month, year)."""
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    return JsonResponse({'status': 'ok', **services.get_channel_quick_payload(month, year)})


@group_required('realise_admin', json_response=True)
@require_http_methods(['POST'])
def api_save_channel_quick_targets(request):
    """Save channel-level targets (one total per channel) for a period. Stored as
    state-blank TargetNodes that survive per-product saves and drive the dashboard's
    channel TGT-L. Admin only."""
    body = _parse_body(request)
    month, year = body.get('month'), body.get('year')
    items = body.get('targets', [])
    if not month or not year or not isinstance(items, list):
        return JsonResponse({'status': 'error', 'error': 'month, year, targets[] required'}, status=400)
    try:
        month, year = int(month), int(year)
    except (ValueError, TypeError):
        return JsonResponse({'status': 'error', 'error': 'month and year must be integers'}, status=400)
    saved = services.save_channel_node_targets(month, year, items, request.user)
    return JsonResponse({'status': 'ok', 'saved': saved})


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_product_actuals(request):
    """Previous-month actual sale per product for one (channel, state) — feeds the
    'last month sold' reference on each product card in the target editor."""
    channel = request.GET.get('channel', '')
    state = request.GET.get('state', '')
    try:
        month = int(request.GET.get('month', datetime.now().month))
        year = int(request.GET.get('year', datetime.now().year))
    except (ValueError, TypeError):
        month, year = datetime.now().month, datetime.now().year
    if not channel:
        return JsonResponse({'status': 'error', 'error': 'channel required'}, status=400)
    return JsonResponse({'status': 'ok', **services.get_product_actuals_payload(channel, state, month, year)})


@never_cache
@group_required(*REALISE_GROUPS, json_response=False)
@require_http_methods(['GET'])
def person_targets_page(request):
    """Person Mapping + Targets page (the React 'Persons Frontend'). Replaces the old
    hierarchical Update Targets editor — reached from the dashboard's Update Targets
    button. Hosts the React app in a same-origin iframe so its global CSS can't touch
    the dashboard chrome. Person reassignment → TerritoryMapping; single target per
    territory → TargetNode."""
    return render(request, 'realise/person_targets.html', {'sidebar_active': 'realise'})


@xframe_options_sameorigin
@never_cache
@group_required(*REALISE_GROUPS, json_response=False)
@require_http_methods(['GET'])
def person_targets_embed(request):
    """Standalone React app (Person Mapping & Targets) shown inside the iframe.
    Marked same-origin-frameable because Django's default X-Frame-Options is DENY,
    which would otherwise block our own iframe."""
    import os
    from django.conf import settings
    from django.middleware.csrf import get_token
    from django.urls import reverse
    from core.context_processors import derive_realise_profile
    role, _, _ = derive_realise_profile(request.user)
    now = datetime.now()

    # Cache-bust the iframe's JS/CSS by their file mtime, so a recompiled bundle is
    # picked up immediately (browsers cache iframe sub-resources aggressively).
    _sdir = os.path.join(str(settings.BASE_DIR), 'realise', 'static', 'realise')
    try:
        asset_ver = int(max(os.path.getmtime(os.path.join(_sdir, 'person_targets.js')),
                            os.path.getmtime(os.path.join(_sdir, 'person_targets.css'))))
    except OSError:
        asset_ver = 1
    boot = {
        'csrf': get_token(request),
        'isAdmin': role == 'admin',
        'month': now.month,
        'year': now.year,
        'monthOptions': list(enumerate(services.MONTHS_ORDER, start=1)),
        'yearOptions': list(range(now.year + 1, now.year - 5, -1)),
        'dashboardUrl': reverse('realise:dashboard') + '#slide2',
        'urls': {
            'map':                reverse('realise:api_territory_map'),
            'mapSave':            reverse('realise:api_save_territory_map'),
            'refresh':            reverse('realise:api_refresh_territory_map'),
            'productTargets':     reverse('realise:api_territory_product_targets'),
            'productTargetsSave': reverse('realise:api_save_territory_product_targets'),
            'channelTargets':     reverse('realise:api_channel_quick_targets'),
            'channelTargetsSave': reverse('realise:api_save_channel_quick_targets'),
            'productActuals':     reverse('realise:api_product_actuals'),
        },
    }
    return render(request, 'realise/person_targets_embed.html', {'boot': boot, 'asset_ver': asset_ver})


@group_required('realise_admin', json_response=False)
@require_http_methods(['GET', 'POST'])
def channel_targets_page(request):
    success = request.GET.get('saved') == '1'
    now = datetime.now()
    error_message = ''
    source = request.POST if request.method == 'POST' else request.GET

    try:
        month = int(source.get('month', now.month))
        year = int(source.get('year', now.year))
    except (TypeError, ValueError):
        month, year = now.month, now.year

    valid_hier = {key for key, _ in services.HIER_ORDERS}
    hier_order = source.get('hier_order', 'mg_state_sp')
    if hier_order not in valid_hier:
        hier_order = 'mg_state_sp'

    hier_filters = {
        'main_group': str(source.get('main_group', '') or '').strip().upper(),
        'state': str(source.get('state', '') or '').strip().upper(),
        'sales_person': str(source.get('sales_person', '') or '').strip().upper(),
    }

    product_segment = services._norm_segment(source.get('segment', ''))

    if request.method == 'POST':
        form_mode = request.POST.get('form_mode', 'segment')
        if form_mode == 'hier':
            keys = request.POST.getlist('node_key')
            vals = request.POST.getlist('node_val')
            reals = request.POST.getlist('node_realise')
            triples = [(keys[i], vals[i] if i < len(vals) else '0',
                        reals[i] if i < len(reals) else '0') for i in range(len(keys))]
            try:
                services.save_hier_targets(month, year, triples, product_segment)
                query = f'hier_order={hier_order}&month={month}&year={year}&saved=1'
                if product_segment:
                    query += f'&segment={product_segment}'
                for key, value in hier_filters.items():
                    if value:
                        query += f'&{key}={value}'
                return redirect(f'/realise/targets/?{query}#hier')
            except ValueError as exc:
                error_message = str(exc)

    master_rows = services.get_territory_master_rows()
    hier_filter_options = services.get_hier_filter_options(master_rows)
    saved_nodes = services.TargetNode.objects.filter(month=month, year=year)
    if product_segment:
        saved_nodes = saved_nodes.filter(segment=product_segment)
    saved_target_map = {
        f'{node.main_group}|{node.state}|{node.sales_person}': {
            'ltrs': float(node.target_ltrs or 0),
            'realise': float(node.target_realise or 0),
        }
        for node in saved_nodes
    }

    context = {
        'sidebar_active': 'realise',
        'target_month': month,
        'target_year': year,
        'product_segment': product_segment,
        'hier_order': hier_order,
        'hier_order_options': services.HIER_ORDERS,
        'hier_filters': hier_filters,
        'hier_filter_options': hier_filter_options,
        'hier_rows': services.get_hier_rows(hier_order, month, year, master_rows, hier_filters, product_segment),
        'hier_master_rows': master_rows,
        'saved_target_map': saved_target_map,
        'month_options': list(enumerate(services.MONTHS_ORDER, start=1)),
        'year_options': list(range(year + 1, year - 5, -1)),
        'save_success': success,
        'error_message': error_message,
    }
    return render(request, 'realise/channel_targets.html', context)


@group_required(*REALISE_GROUPS, json_response=False)
@require_http_methods(['GET'])
def channel_detail_placeholder(request, group):
    return render(request, 'realise/channel_detail_placeholder.html', {
        'sidebar_active': 'realise',
        'channel_group': str(group or '').upper(),
    })


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['GET'])
def api_export_raw_csv(request):
    rows    = _raw_cache.get('rows') or []
    columns = _raw_cache.get('columns') or []
    if not rows or not columns:
        return JsonResponse({'error': 'No data — click Fetch Data first'}, status=400)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for row in rows:
        values = []
        for col in columns:
            val = row.get(col, '')
            if hasattr(val, 'isoformat'):
                val = val.isoformat()
            values.append(val)
        writer.writerow(values)

    start = _raw_cache.get('start', 'from')
    end   = _raw_cache.get('end', 'to')
    filename = f'Sales_RAW_{start}_{end}.csv'
    response = HttpResponse(buf.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _raw_sheet_rows(columns, rows):
    sheet = [[{'value': col, 'style': 1} for col in columns]]
    for row in rows:
        values = []
        for col in columns:
            val = row.get(col, '')
            if hasattr(val, 'isoformat'):
                val = val.isoformat()
            values.append(val)
        sheet.append(values)
    return sheet


@group_required(*REALISE_GROUPS, json_response=True)
@require_http_methods(['POST'])
def api_export_excel(request):
    rows = _raw_cache.get('rows') or []
    columns = _raw_cache.get('columns') or []
    if not rows or not columns:
        return JsonResponse({'error': 'No data - click Fetch Data first'}, status=400)

    body = _parse_body(request)
    layout_rows = body.get('layout_rows') or []
    start = _raw_cache.get('start', 'from')
    end = _raw_cache.get('end', 'to')
    content = build_workbook([
        ('Layout', layout_rows or [['Realise Dashboard']]),
        ('Raw Data', _raw_sheet_rows(columns, rows)),
    ])
    filename = f'Realise_Export_{start}_{end}.xlsx'
    response = HttpResponse(
        content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
