"""Excel export for the Jivo Wellness–Mart reconciliation — the PO chains with their five
document totals, each paired with its underlying SAP document numbers, plus the status."""

from .reconciliation import get_reconciliation


def build_reconciliation_xlsx(date_from=None, date_to=None, only="broken", schema="oil"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    data = get_reconciliation(date_from=date_from, date_to=date_to, schema=schema)
    chains = data.get("chains", [])
    if only == "broken":
        chains = [c for c in chains if c.get("status") != "MATCHED"]
    company = "Beverages" if schema == "beverages" else "Oil"

    head_font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    head_fill = PatternFill("solid", fgColor="FF1E293B")
    bold = Font(name="Calibri", size=10, bold=True)
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    status_fill = {
        "MISMATCH": PatternFill("solid", fgColor="FFFECACA"),
        "INCOMPLETE": PatternFill("solid", fgColor="FFFEF3C7"),
        "MATCHED": PatternFill("solid", fgColor="FFDCFCE7"),
    }

    wb = Workbook()

    # ── Chains sheet ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Reconciliation"
    title = "Jivo Wellness–Mart Reconciliation (%s)   %s → %s   (tolerance ₹%s)" % (
        company, data.get("date_from"), data.get("date_to"), data.get("tolerance"))
    ws.cell(row=1, column=1, value=title).font = Font(name="Calibri", size=12, bold=True)

    # (header, chain-key, kind). kind drives formatting: money = right-aligned thousands,
    # docs = the underlying SAP document numbers joined, status = coloured badge. Each
    # amount is paired with its Doc# column so every figure is traceable to its documents.
    def money(v):
        return None if v is None else round(float(v), 2)

    def docnums(docs):
        return ", ".join(str(d["num"]) for d in (docs or [])) or None

    def docdates(docs):
        # Dates in the same order as the Doc# column, so number ↔ date line up per partial.
        return ", ".join(str(d.get("date") or "?") for d in (docs or [])) or None

    # Each node gets an amount, a Doc# column, and a matching Date column.
    COLS = [
        ("PO No", "po", "text"), ("PO Date", "po_date", "text"), ("Vendor", "vendor", "text"),
        ("Status", "status", "status"), ("Detail", "detail", "text"),
        ("PO Total", "po_total", "money"),
        ("SO", "so", "money"), ("SO Doc#", "so_docs", "docs"), ("SO Date", "so_docs", "dates"),
        ("GRPO", "grpo", "money"), ("GRPO Doc#", "grpo_docs", "docs"), ("GRPO Date", "grpo_docs", "dates"),
        ("A/P", "ap", "money"), ("A/P Doc#", "ap_docs", "docs"), ("A/P Date", "ap_docs", "dates"),
        ("A/R", "ar", "money"), ("A/R Doc#", "ar_docs", "docs"), ("A/R Date", "ar_docs", "dates"),
        ("Delivery", "delivery", "money"), ("Delivery Doc#", "delivery_docs", "docs"), ("Delivery Date", "delivery_docs", "dates"),
    ]

    for i, (h, _k, kind) in enumerate(COLS, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = head_font; c.fill = head_fill
        c.alignment = right if kind == "money" else left

    r = 4
    for ch in chains:
        for i, (_h, key, kind) in enumerate(COLS, start=1):
            if kind == "money":
                v = money(ch.get(key))
            elif kind == "docs":
                v = docnums(ch.get(key))
            elif kind == "dates":
                v = docdates(ch.get(key))
            else:
                v = ch.get(key) or ""
            cell = ws.cell(row=r, column=i, value=v)
            cell.alignment = right if kind == "money" else left
            if kind == "money" and v is not None:
                cell.number_format = '#,##0'
            if kind == "status":
                cell.fill = status_fill.get(ch["status"], status_fill["INCOMPLETE"]); cell.font = bold
        r += 1

    widths = [13, 12, 30, 12, 24, 14,
              13, 22, 22, 13, 22, 22, 13, 20, 22, 13, 24, 24, 12, 20, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B4"   # keep PO No + header rows visible on this wide sheet

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
